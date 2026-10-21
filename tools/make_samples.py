# -*- coding: utf-8 -*-
"""生成测试用 PPTX / PDF 样本，并验证整条转换链路。

用法：
    .venv\\Scripts\\python.exe tools\\make_samples.py
    .venv\\Scripts\\python.exe tools\\test_pipeline.py
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SAMPLES = ROOT / "samples"


def make_pptx() -> Path:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    SAMPLES.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]
    colors = [(0x2A, 0x6D, 0xF4), (0x12, 0xB7, 0x6A), (0xF0, 0x44, 0x38),
              (0xF7, 0x90, 0x09), (0x7C, 0x3A, 0xED)]

    for i in range(5):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.6))
        tf = box.text_frame
        tf.text = f"第 {i + 1} 页 · 测试幻灯片"
        tf.paragraphs[0].font.size = Pt(54)
        tf.paragraphs[0].font.bold = True
        r, g, b = colors[i]
        tf.paragraphs[0].font.color.rgb = RGBColor(r, g, b)

        sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.3), Inches(11.7), Inches(0.9))
        sub.text_frame.text = "Slide → Image pipeline verification 中文渲染检查"
        sub.text_frame.paragraphs[0].font.size = Pt(24)

        bar = slide.shapes.add_shape(1, Inches(0.8), Inches(6.2),
                                     Inches(11.7), Inches(0.25))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor(r, g, b)
        bar.line.fill.background()

    out = SAMPLES / "测试演示.pptx"
    prs.save(str(out))
    return out


def make_pdf() -> Path:
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz

    SAMPLES.mkdir(parents=True, exist_ok=True)
    doc = fitz.open()
    for i in range(3):
        page = doc.new_page(width=595, height=842)  # A4
        page.insert_text((72, 120), f"PDF Page {i + 1}", fontsize=40,
                         color=(0.16, 0.43, 0.96))
        page.insert_text((72, 190), "中文渲染检查 · 幻灯片转图片工具", fontsize=18,
                         fontname="china-s", color=(0.1, 0.1, 0.1))
        page.draw_rect(fitz.Rect(72, 240, 523, 260), color=None,
                       fill=(0.16, 0.43, 0.96))
    out = SAMPLES / "测试文档.pdf"
    doc.save(str(out))
    doc.close()
    return out


if __name__ == "__main__":
    print("生成样例：", make_pptx())
    print("生成样例：", make_pdf())
